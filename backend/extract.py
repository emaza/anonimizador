"""Extracción de texto 100% local para .txt, .md, .docx y .pdf.

Nada se escribe en disco: los bytes llegan en memoria y se descartan al terminar.
No se hace OCR: si un PDF no tiene capa de texto, se avisa explícitamente.
"""
from __future__ import annotations

import io


class ExtractionError(Exception):
    """Error controlado de extracción, con mensaje apto para mostrar al usuario."""


def extract_text(filename: str, data: bytes) -> str:
    """Devuelve el texto plano de un archivo soportado.

    Lanza ExtractionError con un mensaje en español si el formato no se soporta
    o si el PDF es una imagen escaneada sin texto extraíble.
    """
    name = (filename or "").lower()

    if name.endswith((".txt", ".md")):
        return _decode_text(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".pdf"):
        return _extract_pdf(data)
    if name.endswith(".doc"):
        raise ExtractionError(
            "El formato .doc antiguo no está soportado. Guárdalo como .docx y vuelve a intentarlo."
        )
    raise ExtractionError("Formato no soportado. Usa .txt, .md, .docx o .pdf.")


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    # latin-1 no falla nunca, pero por si acaso:
    return data.decode("utf-8", errors="replace")


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(
            "Falta la librería python-docx. Instala las dependencias (ver README)."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("No se pudo leer el .docx (¿archivo dañado?).") from exc

    parts: list[str] = [p.text for p in document.paragraphs]

    # Texto dentro de tablas (python-docx no lo incluye en .paragraphs).
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)

    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError("El documento .docx no contiene texto.")
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(
            "Falta la librería pdfplumber. Instala las dependencias (ver README)."
        ) from exc

    try:
        pages_text: list[str] = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages_text.append(page_text)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("No se pudo leer el PDF (¿archivo dañado?).") from exc

    text = "\n\n".join(pages_text).strip()
    if not text:
        raise ExtractionError(
            "Este PDF no tiene texto extraíble: parece escaneado (imagen). "
            "No se hace OCR, así que tendrás que pegar el texto manualmente."
        )
    return text
